from io import BytesIO

from PIL import Image
from pptx import Presentation


def test_create_edit_validate_export_pipeline(client):
    presentation = client.post("/api/v1/presentations", json={"title": "Pipeline"}).json()
    slide = client.post(f"/api/v1/presentations/{presentation['id']}/slides", json={}).json()
    image_bytes = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(image_bytes, "PNG")
    asset = client.post(f"/api/v1/presentations/{presentation['id']}/assets", files={"file": ("blue.png", image_bytes.getvalue(), "image/png")}).json()
    updated = client.patch(f"/api/v1/slides/{slide['id']}", json={"elements": [{"type": "text", "x": 1, "y": 1, "width": 5, "height": 1, "text": "Build with IR"}, {"type": "shape", "x": 1, "y": 2, "width": 3, "height": 1, "shape_type": "rounded_rectangle"}, {"type": "image", "x": 5, "y": 2, "width": 1, "height": 1, "source": asset["source"]}]} )
    assert updated.status_code == 200
    reloaded = client.get(f"/api/v1/presentations/{presentation['id']}")
    assert reloaded.json()["slides"][0]["elements"][0]["text"] == "Build with IR"
    second_slide = client.post(f"/api/v1/presentations/{presentation['id']}/slides", json={"elements": [{"type": "text", "x": 1, "y": 1, "width": 4, "height": 1, "text": "Second slide"}]}).json()
    assert second_slide["order"] == 1
    assert client.post(f"/api/v1/presentations/{presentation['id']}/validate").status_code == 200
    export = client.post(f"/api/v1/presentations/{presentation['id']}/export")
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert export.content.startswith(b"PK")
    deck = Presentation(BytesIO(export.content))
    assert len(deck.slides) == 2
    assert any(shape.has_text_frame and "Build with IR" in shape.text for shape in deck.slides[0].shapes)
    assert len(deck.slides[0].shapes) == 3
