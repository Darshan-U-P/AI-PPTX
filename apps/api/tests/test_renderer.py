from pathlib import Path

from PIL import Image
from pptx import Presentation

from app.schemas.presentation import PresentationIR
from app.services.pptx_renderer import PPTXRenderer


def render(ir):
    return Presentation(__import__("io").BytesIO(PPTXRenderer().render(ir)))


def test_render_empty_presentation():
    assert len(render(PresentationIR(metadata={"title": "Empty"})).slides) == 0


def test_render_text_shape_image_and_multiple_slides(tmp_path):
    image_path = tmp_path / "image.png"
    Image.new("RGB", (5, 5), "blue").save(image_path)
    ir = PresentationIR(metadata={"title": "Deck"}, slides=[
        {"elements": [{"type": "text", "x": 1, "y": 1, "width": 3, "height": 1, "text": "Hello"}, {"type": "shape", "x": 1, "y": 2, "width": 2, "height": 1, "shape_type": "rectangle"}, {"type": "image", "x": 4, "y": 1, "width": 1, "height": 1, "source": str(image_path)}]},
        {"elements": [{"type": "line", "x": 1, "y": 1, "width": 2, "height": 2}]}
    ])
    deck = render(ir)
    assert len(deck.slides) == 2
    assert len(deck.slides[0].shapes) == 3
    assert deck.slides[0].shapes[0].has_text_frame

