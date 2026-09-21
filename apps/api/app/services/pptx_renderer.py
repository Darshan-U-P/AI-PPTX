from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from app.core.config import assets_path
from app.schemas.presentation import Element, GroupElement, ImageElement, LineElement, PresentationIR, ShapeElement, TextElement
from app.services.renderer import Renderer


def _rgb(value: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value.removeprefix("#").upper())


class PPTXRenderer(Renderer):
    """Deterministically turns validated Presentation IR into editable PPTX objects."""

    def render(self, presentation: PresentationIR) -> bytes:
        deck = Presentation()
        deck.slide_width = Inches(presentation.dimensions.width)
        deck.slide_height = Inches(presentation.dimensions.height)
        blank = deck.slide_layouts[6]
        for ir_slide in sorted(presentation.slides, key=lambda item: item.order):
            slide = deck.slides.add_slide(blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(ir_slide.background)
            for element in sorted(ir_slide.elements, key=lambda item: item.z_index):
                self._add_element(slide, element)
        output = BytesIO()
        deck.save(output)
        return output.getvalue()

    def _add_element(self, slide, element):
        if isinstance(element, TextElement):
            shape = slide.shapes.add_textbox(Inches(element.x), Inches(element.y), Inches(element.width), Inches(element.height))
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.text = element.text
            paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[element.alignment]
            run = paragraph.runs[0]
            run.font.name = element.font_family
            run.font.size = Pt(element.font_size)
            run.font.bold = element.font_weight >= 600
            run.font.color.rgb = _rgb(element.color)
            shape.text_frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[element.vertical_alignment]
            paragraph.space_after = Pt(0)
            shape.rotation = element.rotation
        elif isinstance(element, ShapeElement):
            shape_type = {"rectangle": MSO_AUTO_SHAPE_TYPE.RECTANGLE, "ellipse": MSO_AUTO_SHAPE_TYPE.OVAL, "rounded_rectangle": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE}[element.shape_type]
            shape = slide.shapes.add_shape(shape_type, Inches(element.x), Inches(element.y), Inches(element.width), Inches(element.height))
            shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(element.fill)
            shape.line.color.rgb = _rgb(element.stroke)
            shape.line.width = Pt(element.stroke_width)
            shape.rotation = element.rotation
        elif isinstance(element, LineElement):
            line = slide.shapes.add_connector(1, Inches(element.x), Inches(element.y), Inches(element.x + element.width), Inches(element.y + element.height))
            line.line.color.rgb = _rgb(element.stroke)
            line.line.width = Pt(element.stroke_width)
        elif isinstance(element, ImageElement):
            parsed = urlparse(element.source)
            if parsed.scheme in {"http", "https"}:
                return  # Remote downloading is intentionally deferred; assets must be locally resolved first.
            path = self._image_path(element.source)
            if path.is_file():
                slide.shapes.add_picture(str(path), Inches(element.x), Inches(element.y), width=Inches(element.width), height=Inches(element.height))
        elif isinstance(element, GroupElement):
            for child in element.children:
                # Groups are a semantic IR container in Phase 0; child objects remain editable PPTX shapes.
                from pydantic import TypeAdapter
                self._add_element(slide, TypeAdapter(Element).validate_python(child))

    @staticmethod
    def _image_path(source: str) -> Path:
        if source.startswith("asset://"):
            return assets_path() / Path(source.removeprefix("asset://")).name
        return Path(source)
